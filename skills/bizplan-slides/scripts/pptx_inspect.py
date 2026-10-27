#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""발표 템플릿(.pptx/.potx)을 분석해 채우기 계획에 필요한 정보를 출력한다.

출력: 슬라이드 크기, 테마 색·글꼴, 레이아웃별 자리표시자, 기존 슬라이드의 도형·글자, 경고.
    python3 -B pptx_inspect.py 템플릿.pptx [--json 프로필.json]
경고 예: 16:9 슬라이드에 4:3 크기 자리표시자, 동아시아(한글) 글꼴 미지정, 레이아웃 없이 샘플 슬라이드만 있는 템플릿.
"""
import argparse
import io
import json
import sys
import zipfile
import re

from pptx import Presentation
from pptx.util import Emu

NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
DEFAULT_LAYOUTS = {"Title Slide", "Title and Content", "Section Header", "Two Content", "Comparison",
                   "Title Only", "Blank", "Content with Caption", "Picture with Caption",
                   "Title and Vertical Text", "Vertical Title and Text"}


def open_presentation(path):
    """.pptx는 그대로, .potx(서식 파일)는 내용 형식만 프레젠테이션으로 바꿔 메모리에서 연다."""
    with zipfile.ZipFile(path) as z:
        ct = z.read("[Content_Types].xml").decode("utf-8")
        if "presentationml.template.main+xml" not in ct:
            return Presentation(path)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
            for item in z.infolist():
                data = z.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = ct.replace("presentationml.template.main+xml",
                                      "presentationml.presentation.main+xml").encode("utf-8")
                out.writestr(item, data)
    buf.seek(0)
    return Presentation(buf)


def _box(shape, sw, sh):
    """도형 위치·크기를 슬라이드 대비 %로."""
    if shape.left is None or shape.width is None:
        return None
    return [round(shape.left / sw * 100, 1), round(shape.top / sh * 100, 1),
            round(shape.width / sw * 100, 1), round(shape.height / sh * 100, 1)]


def _theme(path):
    """테마 XML에서 색(dk1·lt1·accent1~6)과 글꼴(major/minor의 latin·ea)을 읽는다."""
    colors, fonts = {}, {}
    with zipfile.ZipFile(path) as z:
        names = sorted(n for n in z.namelist() if re.match(r"ppt/theme/theme\d+\.xml$", n))
        if not names:
            return colors, fonts
        xml = z.read(names[0]).decode("utf-8", "ignore")
    for key in ["dk1", "lt1", "dk2", "lt2"] + [f"accent{i}" for i in range(1, 7)]:
        m = re.search(rf"<a:{key}>.*?(?:srgbClr val=\"([0-9A-Fa-f]{{6}})\"|lastClr=\"([0-9A-Fa-f]{{6}})\")", xml, re.S)
        if m:
            colors[key] = "#" + (m.group(1) or m.group(2)).upper()
    for kind in ("major", "minor"):
        m = re.search(rf"<a:{kind}Font>(.*?)</a:{kind}Font>", xml, re.S)
        if m:
            block = m.group(1)
            latin = re.search(r'<a:latin typeface="([^"]*)"', block)
            ea = re.search(r'<a:ea typeface="([^"]*)"', block)
            hang = re.search(r'<a:font script="Hang" typeface="([^"]*)"', block)
            fonts[kind] = {"latin": latin.group(1) if latin else "",
                           "ea": ea.group(1) if ea else "",
                           "Hang": hang.group(1) if hang else ""}
    return colors, fonts


OFFICE_ACCENT1 = {"#4F81BD", "#4472C4", "#156082", "#0F9ED5"}   # Office 기본 테마 accent1


def _fill(sp):
    """도형의 단색 채우기 색(#RRGGBB). 없거나 테마 색이면 None."""
    spPr = sp._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
    el = spPr.find(NS_A + "solidFill") if spPr is not None else None
    if el is None:
        return None
    rgb = el.find(NS_A + "srgbClr")
    return "#" + rgb.get("val").upper() if rgb is not None else None


def _deco(shapes, sw, sh):
    """자리표시자가 아닌 도형(띠·로고·바닥글 글상자 등)을 나열한다."""
    out = []
    for sp in shapes:
        if sp.is_placeholder:
            continue
        item = {"name": sp.name, "kind": sp.shape_type and str(sp.shape_type).split(".")[-1].split(" ")[0],
                "box_pct": _box(sp, sw, sh)}
        color = _fill(sp)
        if color:
            item["fill"] = color
        if sp.has_text_frame and sp.text_frame.text.strip():
            item["text"] = sp.text_frame.text.strip()[:80]
            run_color = sp._element.find(".//" + NS_A + "rPr/" + NS_A + "solidFill/" + NS_A + "srgbClr")
            if run_color is not None:
                item["text_color"] = "#" + run_color.get("val").upper()
        out.append(item)
    return out


def inspect(path):
    prs = open_presentation(path)
    sw, sh = prs.slide_width, prs.slide_height
    ratio = sw / sh
    profile = {
        "file": path,
        "slide_size_in": [round(Emu(sw).inches, 3), round(Emu(sh).inches, 3)],
        "aspect": "16:9" if abs(ratio - 16 / 9) < 0.02 else "4:3" if abs(ratio - 4 / 3) < 0.02 else f"{ratio:.3f}",
        "layouts": [], "slides": [], "warnings": [],
    }
    profile["theme_colors"], profile["theme_fonts"] = _theme(path)
    profile["master_shapes"] = _deco(prs.slide_master.shapes, sw, sh)
    profile["layout_shapes"] = {l.name: _deco(l.shapes, sw, sh) for l in prs.slide_layouts if _deco(l.shapes, sw, sh)}

    narrow = []
    for li, layout in enumerate(prs.slide_layouts):
        phs = []
        for ph in layout.placeholders:
            phs.append({"idx": ph.placeholder_format.idx, "type": str(ph.placeholder_format.type).split(".")[-1].split(" ")[0],
                        "name": ph.name, "box_pct": _box(ph, sw, sh)})
        profile["layouts"].append({"index": li, "name": layout.name, "placeholders": phs})
        # 16:9인데 본문 자리표시자들의 오른쪽 끝이 80%에 못 미치면 4:3에서 늘리지 않은 레이아웃일 가능성
        boxes = [p["box_pct"] for p in phs if p["box_pct"] and p["type"] not in ("DATE", "FOOTER", "SLIDE_NUMBER")]
        # 4:3(10in) 기준 배치를 16:9(13.33in)에 그대로 두면 '왼쪽 여백 + 오른쪽 끝'이 약 75%가 된다.
        # 가운데 정렬로 좁게 둔 디자인(합이 약 100%)은 잔재가 아니다.
        if profile["aspect"] == "16:9" and boxes:
            right = max(b[0] + b[2] for b in boxes)
            left = min(b[0] for b in boxes)
            if right < 80 and left + right < 85:
                narrow.append(layout.name)
    if narrow:
        profile["warnings"].append(
            f"16:9 슬라이드인데 자리표시자가 왼쪽 80% 안에 몰린 레이아웃: {', '.join(narrow)} (4:3 잔재 의심) — "
            "pptx_build.py --stretch-placeholders로 가로를 맞추거나 템플릿을 수정")

    for si, slide in enumerate(prs.slides):
        shapes = []
        for sp in slide.shapes:
            item = {"id": sp.shape_id, "name": sp.name, "kind": sp.shape_type and str(sp.shape_type).split(".")[-1].split(" ")[0],
                    "box_pct": _box(sp, sw, sh)}
            if sp.is_placeholder:
                item["placeholder_idx"] = sp.placeholder_format.idx
            if sp.has_text_frame and sp.text_frame.text.strip():
                item["text"] = sp.text_frame.text.strip()[:120]
            if getattr(sp, "has_table", False) and sp.has_table:
                t = sp.table
                item["table"] = [len(t.rows), len(t.columns)]
            shapes.append(item)
        profile["slides"].append({"index": si + 1, "layout": slide.slide_layout.name, "shapes": shapes})

    # 실제 디자인 색: 마스터·레이아웃·샘플 슬라이드 도형의 채우기·글자색(테마 accent가 Office 기본값이면 이쪽이 진짜)
    colors = []
    for item in profile["master_shapes"] + [i for v in profile["layout_shapes"].values() for i in v]:
        for key in ("fill", "text_color"):
            if item.get(key) and item[key] not in colors and item[key] not in ("#FFFFFF", "#000000"):
                colors.append(item[key])
    for slide in prs.slides:
        for item in _deco(slide.shapes, sw, sh):
            if item.get("fill") and item["fill"] not in colors and item["fill"] not in ("#FFFFFF", "#000000"):
                colors.append(item["fill"])
    profile["design_colors"] = colors
    if profile["theme_colors"].get("accent1") in OFFICE_ACCENT1 and colors:
        profile["warnings"].append(
            f"테마 색이 Office 기본값이다 — 차트·강조색은 실제 디자인 색 {', '.join(colors[:3])}을 쓴다(use_palette)")

    # 디자인이 마스터·레이아웃에 없고 샘플 슬라이드에만 있을 때만 복제 방식을 권한다
    custom = [l["name"] for l in profile["layouts"] if l["name"] not in DEFAULT_LAYOUTS]
    slide_deco = any(_deco(s.shapes, sw, sh) for s in prs.slides)
    if profile["slides"] and not custom and not profile["master_shapes"] and not profile["layout_shapes"] and slide_deco:
        profile["warnings"].append(
            "마스터·레이아웃에는 디자인 도형이 없고 샘플 슬라이드에만 있다 — "
            "spec에서 from_slide(샘플 슬라이드 복제 + 글 교체, 필요하면 add로 그림·표 추가)를 우선 사용")
    fonts = profile["theme_fonts"]
    if fonts and not any(f.get("ea") or f.get("Hang") for f in fonts.values()):
        profile["warnings"].append("테마에 한글(동아시아) 글꼴이 없다 — pptx_build.py가 --ea-font(기본 '맑은 고딕')로 지정한다")
    return profile


def _print(profile):
    print(f"파일: {profile['file']}")
    print(f"크기: {profile['slide_size_in'][0]}×{profile['slide_size_in'][1]}in ({profile['aspect']})")
    if profile["theme_colors"]:
        print("테마 색: " + ", ".join(f"{k}={v}" for k, v in profile["theme_colors"].items()))
    if profile["theme_fonts"]:
        print("테마 글꼴: " + "; ".join(f"{k}: latin={v['latin']} ea={v['ea'] or v['Hang'] or '-'}"
                                      for k, v in profile["theme_fonts"].items()))
    if profile.get("design_colors"):
        print("디자인 색(도형 채우기·글자): " + ", ".join(profile["design_colors"]))
    if profile.get("master_shapes"):
        print("\n[마스터 도형 — 모든 장에 나타남, 지우거나 옮기지 않는다]")
        for it in profile["master_shapes"]:
            extra = (f" 채우기 {it['fill']}" if it.get("fill") else "") + (f" \"{it['text']}\"" if it.get("text") else "")
            print(f"  - {it['name']} ({it['kind']}) {it['box_pct']}{extra}")
    for lname, items in profile.get("layout_shapes", {}).items():
        print(f"\n[레이아웃 '{lname}' 고유 도형]")
        for it in items:
            extra = (f" 채우기 {it['fill']}" if it.get("fill") else "") + (f" \"{it['text']}\"" if it.get("text") else "")
            print(f"  - {it['name']} ({it['kind']}) {it['box_pct']}{extra}")
    print("\n[레이아웃]")
    for l in profile["layouts"]:
        phs = ", ".join(f"{p['idx']}:{p['type']}{p['box_pct']}" for p in l["placeholders"])
        print(f"  {l['index']:>2} {l['name']} — {phs or '자리표시자 없음'}")
    print("\n[기존 슬라이드]")
    for s in profile["slides"]:
        print(f"  #{s['index']} (레이아웃: {s['layout']})")
        for sp in s["shapes"]:
            extra = f" \"{sp['text']}\"" if sp.get("text") else ""
            if sp.get("table"):
                extra += f" 표{sp['table'][0]}×{sp['table'][1]}"
            ph = f" ph{sp['placeholder_idx']}" if "placeholder_idx" in sp else ""
            print(f"     - [{sp['id']}] {sp['name']} ({sp['kind']}{ph}) {sp['box_pct']}{extra}")
    if profile["warnings"]:
        print("\n[경고]")
        for w in profile["warnings"]:
            print("  ! " + w)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("template")
    ap.add_argument("--json", help="프로필을 JSON으로 저장할 경로")
    args = ap.parse_args()
    profile = inspect(args.template)
    _print(profile)
    if args.json:
        with open(args.json, "w", encoding="utf-8") as f:
            json.dump(profile, f, ensure_ascii=False, indent=2)
        print(f"\nJSON 저장: {args.json}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
