#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""발표 템플릿(.pptx/.potx)의 디자인을 유지한 채 슬라이드 내용을 채워 .pptx를 만든다.

    python3 -B pptx_build.py --spec deck.json --out 발표.pptx [--template 템플릿.pptx]
                             [--ea-font "맑은 고딕"] [--stretch-placeholders] [--keep-template-slides]

spec(JSON) 형식 — 슬라이드마다 둘 중 하나:
  레이아웃 채우기  {"layout": "Title and Content", "title": "...", "bullets": ["문장", ["하위1", "하위2"]],
                   "image": "차트.png", "caption": "결론 한 줄", "table": {"header": [...], "rows": [[...]], "highlight_row": 2},
                   "left": {...}, "right": {...}, "subtitle": "...", "source": "출처", "notes": "발표 대본"}
  샘플 복제·교체  {"from_slide": 3, "replace": {"도형 이름 또는 #도형id": "새 글" | ["불릿", ...] | {"image": "a.png"}
                   | {"table": {...}}}, "notes": "..."}
레이아웃은 이름(부분 일치) 또는 번호. 생략하면 내용에 맞춰 고른다(이미지·표만 → Title Only, 그 외 → Title and Content).
템플릿에 원래 있던 샘플 슬라이드는 기본으로 지운다(--keep-template-slides로 유지).
글이 상자를 넘칠 것으로 추정되면 글자 크기를 2pt씩 줄이고(본문 최소 14pt, 제목 최소 20pt) 그래도 넘치면 경고한다.
추정은 근사치다 — 반드시 export_pdf.py + preview.py 렌더로 눈으로 확인한다.
"""
import argparse
import copy
import json
import math
import os
import re
import sys

from pptx import Presentation

from pptx_inspect import open_presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}
FOOT_TYPES = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
RID_ATTRS = (qn("r:embed"), qn("r:link"), qn("r:id"))

warnings = []


def warn(slide_no, msg):
    warnings.append(f"#{slide_no}: {msg}")


# ---------------------------------------------------------------- 글꼴·크기

def _ea(run, font):
    """런에 동아시아(한글) 글꼴을 지정한다. rPr 자식 순서(latin→ea→cs→sym…)를 지킨다."""
    if not font:
        return
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        after = [rpr.find(qn(t)) for t in ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")]
        after = [a for a in after if a is not None]
        if after:
            after[0].addprevious(ea)
        else:
            rpr.append(ea)
    ea.set("typeface", font)


def _inherited_size(shape, level=0, title=False):
    """자리표시자 글자 크기(pt)를 레이아웃→마스터 스타일에서 찾는다. 못 찾으면 기본값."""
    lvl = f"a:lvl{level + 1}pPr"
    try:
        if shape.is_placeholder:
            layout_ph = shape.part.slide_layout.placeholders.get(idx=shape.placeholder_format.idx) \
                if hasattr(shape.part, "slide_layout") else None
            for el in [shape._element, layout_ph._element if layout_ph is not None else None]:
                if el is None:
                    continue
                d = el.find(f".//{qn('a:lstStyle')}/{qn(lvl)}/{qn('a:defRPr')}")
                if d is not None and d.get("sz"):
                    return int(d.get("sz")) / 100
            master = shape.part.slide_layout.slide_master
            style = master._element.find(qn("p:txStyles"))
            key = "p:titleStyle" if title else "p:bodyStyle"
            d = style.find(f"{qn(key)}/{qn(lvl)}/{qn('a:defRPr')}") if style is not None else None
            if d is not None and d.get("sz"):
                return int(d.get("sz")) / 100
    except Exception:
        pass
    return 32.0 if title else 20.0


def _em(text):
    """글자 폭을 em 단위로 근사: 한글·전각 1.0, 영숫자 0.6, 공백 0.33."""
    w = 0.0
    for ch in text:
        if ch.isspace():
            w += 0.33
        elif ord(ch) > 0x2E7F:
            w += 1.0
        else:
            w += 0.6
    return w


def _needed_height(paras, size, width_pt):
    """paras=[(글, 수준)] 를 size pt로 썼을 때 필요한 높이(pt) 근사."""
    h = 0.0
    for text, level in paras:
        # 들여쓰기·불릿 몫을 빼고, 대체 글꼴·한영 자간 차이를 감안해 폭을 8% 덜 쓴다고 본다
        avail = max((width_pt - level * size * 1.2 - size * 0.8) * 0.92, size * 4)
        lines = max(1, math.ceil(_em(text) * size / avail))
        h += lines * size * 1.2 + size * 0.25
    return h


def _fit(shape, paras, size, title, slide_no):
    """넘칠 것 같으면 크기를 줄인 값을 돌려준다(줄일 필요 없으면 None)."""
    tf = shape.text_frame
    width = Emu(shape.width).pt - Emu(tf.margin_left).pt - Emu(tf.margin_right).pt
    height = Emu(shape.height).pt - Emu(tf.margin_top).pt - Emu(tf.margin_bottom).pt
    floor = 20 if title else 14
    s = size
    while _needed_height(paras, s, width) > height and s - 2 >= floor:
        s -= 2
    if _needed_height(paras, s, width) > height:
        warn(slide_no, f"'{shape.name}' 글이 넘칠 수 있다({s:.0f}pt에서도) — 문장을 줄이거나 장을 나눈다")
    return s if s != size else None


# ---------------------------------------------------------------- 글 채우기

def _flatten(items, level=0):
    out = []
    for it in items:
        if isinstance(it, list):
            out += _flatten(it, level + 1)
        else:
            out.append((str(it), level))
    return out


def _set_paragraphs(shape, paras, ea_font, slide_no, title=False, keep_format=False):
    """글상자에 문단들을 쓴다. keep_format이면 첫 문단·첫 런의 서식을 복제해 쓴다(샘플 복제용)."""
    tf = shape.text_frame
    proto_p = copy.deepcopy(tf.paragraphs[0]._p) if keep_format else None
    proto_r = None
    if proto_p is not None:
        rs = proto_p.findall(qn("a:r"))
        proto_r = copy.deepcopy(rs[0]) if rs else None
        for child in list(proto_p):
            if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
                proto_p.remove(child)
    size = None
    if keep_format and proto_r is not None:
        rpr = proto_r.find(qn("a:rPr"))
        if rpr is not None and rpr.get("sz"):
            size = int(rpr.get("sz")) / 100
    if size is None:
        size = _inherited_size(shape, 0, title) if shape.is_placeholder else (24.0 if title else 18.0)
    fitted = _fit(shape, paras, size, title, slide_no)
    paras = paras or [("", 0)]

    txBody = tf._txBody
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    for text, level in paras:
        if proto_p is not None:
            p = copy.deepcopy(proto_p)
            txBody.append(p)
            r = copy.deepcopy(proto_r) if proto_r is not None else p.makeelement(qn("a:r"), {})
            t = r.find(qn("a:t"))
            if t is None:
                t = r.makeelement(qn("a:t"), {})
                r.append(t)
            t.text = text
            end = p.find(qn("a:endParaRPr"))
            (end.addprevious(r) if end is not None else p.append(r))
            para = next(pp for pp in tf.paragraphs if pp._p is p)
        else:
            para = tf.add_paragraph()
            # 일반 글상자에는 불릿 서식이 없으므로 기호를 글로 붙인다
            para.text = (("• ", "– ", "· ")[min(level, 2)] + text) if not shape.is_placeholder and not title else text
        if level and shape.is_placeholder:
            para.level = min(level, 4)
        for run in para.runs:
            _ea(run, ea_font)
            if fitted:
                lvl_size = fitted if not level else max(fitted - 2 * level, 12)
                run.font.size = Pt(lvl_size)


# ---------------------------------------------------------------- 그림·표·캡션

def _content_box(slide, prs, title_shape):
    """본문 자리표시자가 없을 때 쓸 영역: 제목 아래 ~ 바닥글 위, 제목과 같은 좌우 여백."""
    sw, sh = prs.slide_width, prs.slide_height
    left = title_shape.left if title_shape is not None else int(sw * 0.05)
    left = max(left, int(sw * 0.04))
    top = (title_shape.top + title_shape.height + int(sh * 0.02)) if title_shape is not None else int(sh * 0.12)
    bottom = _footer_top(slide, prs) - int(sh * 0.06)
    return left, top, sw - left * 2, max(bottom - top, int(sh * 0.3))


def _footer_top(slide, prs):
    tops = [ph.top for ph in slide.slide_layout.placeholders
            if ph.placeholder_format.type in FOOT_TYPES and ph.top is not None]
    return min(tops) if tops else int(prs.slide_height * 0.93)


def _picture(slide, path, box, slide_no):
    if not os.path.exists(path):
        warn(slide_no, f"그림 파일 없음: {path}")
        return None
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    left, top, width, height = box
    scale = min(width / iw, height / ih)
    w, h = int(iw * scale), int(ih * scale)
    return slide.shapes.add_picture(path, left + (width - w) // 2, top + (height - h) // 2, w, h)


def _textbox(slide, box, text, size, ea_font, color=None, align=None):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    tb = slide.shapes.add_textbox(*box)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.text = text
    for p in tf.paragraphs:
        if align == "center":
            p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
            _ea(r, ea_font)
    return tb


def _table(slide, spec, box, ea_font, slide_no):
    header, rows = spec.get("header", []), spec.get("rows", [])
    ncols = len(header) or (len(rows[0]) if rows else 0)
    if not ncols:
        warn(slide_no, "표에 열이 없다")
        return None
    nrows = len(rows) + (1 if header else 0)
    left, top, width, height = box
    row_h = min(int(height / max(nrows, 1)), int(Pt(34)))
    gt = slide.shapes.add_table(nrows, ncols, left, top, width, row_h * nrows)
    table = gt.table
    # 열 너비: 열별 최장 글 길이에 비례(최소 12%)
    lens = [max([_em(str(r[c])) for r in rows if c < len(r)] + [_em(str(header[c])) if header else 1, 2])
            for c in range(ncols)]
    total = sum(lens)
    shares = [max(l / total, 0.12) for l in lens]
    norm = sum(shares)
    for c in range(ncols):
        table.columns[c].width = int(width * shares[c] / norm)
    size = 14 if nrows <= 6 else 12 if nrows <= 9 else 10
    all_rows = ([header] if header else []) + rows
    hl = spec.get("highlight_row")
    for ri, row in enumerate(all_rows):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            cell.text = str(row[ci]) if ci < len(row) else ""
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
                    if (header and ri == 0) or (hl is not None and ri == hl + (1 if header else 0)):
                        r.font.bold = True
                    _ea(r, ea_font)
    return gt


# ---------------------------------------------------------------- 레이아웃 슬라이드

def _pick_layout(prs, spec, slide_no):
    want = spec.get("layout")
    layouts = list(prs.slide_layouts)
    if isinstance(want, int):
        return layouts[want]
    if want:
        for l in layouts:
            if l.name == want:
                return l
        for l in layouts:
            if want.lower() in l.name.lower():
                return l
        warn(slide_no, f"레이아웃 '{want}' 없음 — 내용으로 대신 고름")
    only_visual = ("image" in spec or "table" in spec) and "bullets" not in spec
    pref = ["Title Only", "제목만"] if only_visual else ["Title and Content", "제목 및 내용"]
    for name in pref:
        for l in layouts:
            if name.lower() in l.name.lower():
                return l
    return layouts[min(1, len(layouts) - 1)]


def _fill_region(slide, prs, region, box, ph, ea_font, slide_no):
    """region={bullets|image|table|caption} 을 box(또는 ph 자리표시자)에 채운다."""
    if "bullets" in region and ph is not None:
        _set_paragraphs(ph, _flatten(region["bullets"]), ea_font, slide_no)
        return
    if ph is not None:
        box = (ph.left, ph.top, ph.width, ph.height)
        ph._element.getparent().remove(ph._element)
    if "bullets" in region:
        tb = slide.shapes.add_textbox(*box)
        tb.text_frame.word_wrap = True
        _set_paragraphs(tb, _flatten(region["bullets"]), ea_font, slide_no)
        return
    cap_h = int(prs.slide_height * 0.06) if region.get("caption") else 0
    inner = (box[0], box[1], box[2], box[3] - cap_h)
    if "image" in region:
        _picture(slide, region["image"], inner, slide_no)
    elif "table" in region:
        _table(slide, region["table"], inner, ea_font, slide_no)
    if region.get("caption"):
        _textbox(slide, (box[0], box[1] + box[3] - cap_h, box[2], cap_h), region["caption"], 14, ea_font,
                 color="334155", align="center")


def build_layout_slide(prs, spec, ea_font, slide_no):
    layout = _pick_layout(prs, spec, slide_no)
    slide = prs.slides.add_slide(layout)
    phs = {"title": None, "subtitle": None, "body": [], "picture": None}
    for ph in slide.placeholders:
        t = ph.placeholder_format.type
        if t in TITLE_TYPES:
            phs["title"] = ph
        elif t == PP_PLACEHOLDER.SUBTITLE:
            phs["subtitle"] = ph
        elif t in BODY_TYPES:
            phs["body"].append(ph)
        elif t == PP_PLACEHOLDER.PICTURE:
            phs["picture"] = ph
    phs["body"].sort(key=lambda p: (p.left or 0))
    used = set()
    if spec.get("title") and phs["title"] is not None:
        _set_paragraphs(phs["title"], [(spec["title"], 0)], ea_font, slide_no, title=True)
        used.add(phs["title"].shape_id)
    if spec.get("subtitle"):
        target = phs["subtitle"] or (phs["body"][0] if phs["body"] else None)
        if target is not None:
            _set_paragraphs(target, [(spec["subtitle"], 0)], ea_font, slide_no)
            used.add(target.shape_id)
            if target in phs["body"]:
                phs["body"].remove(target)

    regions = []
    if "left" in spec or "right" in spec:
        regions = [spec.get("left", {}), spec.get("right", {})]
    else:
        main = {k: spec[k] for k in ("bullets", "image", "table", "caption") if k in spec}
        if main:
            regions = [main]
    default_box = _content_box(slide, prs, phs["title"])
    bodies = list(phs["body"])
    if phs["picture"] is not None and regions and "image" in regions[0] and "bullets" not in regions[0]:
        bodies = [phs["picture"]] + bodies
    if len(regions) == 2 and len(bodies) < 2:
        l, t, w, h = default_box
        boxes = [(l, t, w // 2 - int(w * 0.02), h), (l + w // 2 + int(w * 0.02), t, w // 2 - int(w * 0.02), h)]
        for region, box in zip(regions, boxes):
            ph = bodies.pop(0) if bodies else None
            if ph is not None:
                used.add(ph.shape_id)
            _fill_region(slide, prs, region, box, ph, ea_font, slide_no)
    else:
        for region in regions:
            ph = bodies.pop(0) if bodies else None
            if ph is not None:
                used.add(ph.shape_id)
            _fill_region(slide, prs, region, default_box, ph, ea_font, slide_no)

    # 안 쓴 빈 자리표시자는 지운다(편집 화면의 '텍스트를 입력하십시오' 방지)
    for ph in list(slide.placeholders):
        if ph.shape_id not in used and not (ph.has_text_frame and ph.text_frame.text.strip()):
            ph._element.getparent().remove(ph._element)
    _source_and_notes(slide, prs, spec, ea_font)
    return slide


def _source_and_notes(slide, prs, spec, ea_font):
    if spec.get("source"):
        sw, sh = prs.slide_width, prs.slide_height
        bottom = _footer_top(slide, prs)
        h = int(sh * 0.045)
        _textbox(slide, (int(sw * 0.04), bottom - h - int(sh * 0.005), int(sw * 0.7), h),
                 spec["source"], 10, ea_font, color="64748B")
    if spec.get("notes"):
        slide.notes_slide.notes_text_frame.text = spec["notes"]


# ---------------------------------------------------------------- 샘플 슬라이드 복제

def clone_slide(prs, src):
    """샘플 슬라이드를 끝에 복제한다(도형·배경·그림 관계 포함)."""
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rel in src.part.rels.values():
        if rel.reltype.endswith(("/slideLayout", "/notesSlide")):
            continue
        if rel.is_external:
            rid_map[rel.rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rel.rId] = new.part.relate_to(rel.target_part, rel.reltype)
    src_csld, new_csld = src._element.cSld, new._element.cSld
    bg = src_csld.find(qn("p:bg"))
    if bg is not None:
        new_csld.insert(0, copy.deepcopy(bg))
    tree = new.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        dup = copy.deepcopy(el)
        for node in dup.iter():
            for attr in RID_ATTRS:
                if node.get(attr) in rid_map:
                    node.set(attr, rid_map[node.get(attr)])
        tree.append(dup)
    return new


def _find_shape(slide, key):
    if isinstance(key, str) and key.startswith("#") and key[1:].isdigit():
        return next((s for s in slide.shapes if s.shape_id == int(key[1:])), None)
    return next((s for s in slide.shapes if s.name == key), None)


def build_cloned_slide(prs, originals, spec, ea_font, slide_no):
    n = spec["from_slide"]
    if not 1 <= n <= len(originals):
        warn(slide_no, f"from_slide {n} 없음(템플릿 슬라이드 {len(originals)}장)")
        return None
    slide = clone_slide(prs, originals[n - 1])
    for key, value in (spec.get("replace") or {}).items():
        shape = _find_shape(slide, key)
        if shape is None:
            warn(slide_no, f"도형 '{key}' 없음 — pptx_inspect.py 출력의 이름/#id를 쓴다")
            continue
        if isinstance(value, dict) and "image" in value:
            box = (shape.left, shape.top, shape.width, shape.height)
            shape._element.getparent().remove(shape._element)
            _picture(slide, value["image"], box, slide_no)
        elif isinstance(value, dict) and "table" in value:
            if not getattr(shape, "has_table", False):
                warn(slide_no, f"'{key}'는 표가 아니다")
                continue
            data = ([value["table"].get("header")] if value["table"].get("header") else []) + value["table"].get("rows", [])
            t = shape.table
            for ri, row in enumerate(data[: len(t.rows)]):
                for ci, v in enumerate(row[: len(t.columns)]):
                    cell = t.cell(ri, ci)
                    paras = cell.text_frame.paragraphs
                    if paras and paras[0].runs:
                        paras[0].runs[0].text = str(v)
                        for r in paras[0].runs[1:]:
                            r._r.getparent().remove(r._r)
                    else:
                        cell.text = str(v)
            if len(data) > len(t.rows):
                warn(slide_no, f"표 '{key}' 행이 모자라 {len(data) - len(t.rows)}행을 버렸다")
        elif shape.has_text_frame:
            paras = _flatten(value) if isinstance(value, list) else [(str(value), 0)]
            is_title = shape.is_placeholder and shape.placeholder_format.type in TITLE_TYPES
            _set_paragraphs(shape, paras, ea_font, slide_no, title=is_title, keep_format=True)
        else:
            warn(slide_no, f"'{key}'에는 글을 넣을 수 없다")
    _source_and_notes(slide, prs, spec, ea_font)
    return slide


# ---------------------------------------------------------------- 4:3 잔재 보정

def stretch_placeholders(prs):
    """16:9 슬라이드에서 4:3 화면 기준으로 놓인(오른쪽 끝이 80% 미만) 마스터·레이아웃 자리표시자를 가로로 늘린다.

    위치를 직접 가진 자리표시자만 모아 한 배율로 늘린다(상속받는 것은 마스터를 따라간다).
    """
    sw = prs.slide_width
    own = []
    for phs in [prs.slide_master.placeholders] + [l.placeholders for l in prs.slide_layouts]:
        for p in phs:
            if p._element.spPr.find(qn("a:xfrm")) is not None and p.width:
                own.append(p)
    if not own:
        return 0
    right = max(p.left + p.width for p in own)
    left_min = min(p.left for p in own)
    if right >= sw * 0.8:
        return 0
    factor = sw / (right + left_min)
    for p in own:
        p.left, p.width = int(p.left * factor), int(p.width * factor)
    return len(own)


# ---------------------------------------------------------------- main

def _theme_has_ea(prs):
    theme_part = None
    for rel in prs.slide_master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            theme_part = rel.target_part
    if theme_part is None:
        return False
    xml = theme_part.blob.decode("utf-8", "ignore")
    return bool(re.search(r'<a:ea typeface="[^"]+"', xml) or 'script="Hang"' in xml)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--spec", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--template", help="spec의 template보다 우선")
    ap.add_argument("--ea-font", help="한글 글꼴(기본: 테마에 없을 때만 '맑은 고딕')")
    ap.add_argument("--stretch-placeholders", action="store_true")
    ap.add_argument("--keep-template-slides", action="store_true")
    args = ap.parse_args()

    with open(args.spec, encoding="utf-8") as f:
        spec = json.load(f)
    base = os.path.dirname(os.path.abspath(args.spec))
    template = args.template or spec.get("template")
    if template and not os.path.isabs(template) and not os.path.exists(template):
        template = os.path.join(base, template)
    prs = open_presentation(template) if template else Presentation()
    ea_font = args.ea_font or spec.get("ea_font") or (None if _theme_has_ea(prs) else "맑은 고딕")
    if args.stretch_placeholders or spec.get("stretch_placeholders"):
        print(f"4:3 잔재 보정: 자리표시자 {stretch_placeholders(prs)}개")

    originals = list(prs.slides)
    for i, s in enumerate(spec.get("slides", []), 1):
        for k in ("image",):
            if k in s and not os.path.isabs(s[k]):
                s[k] = os.path.join(base, s[k])
        for side in ("left", "right"):
            if isinstance(s.get(side), dict) and "image" in s[side] and not os.path.isabs(s[side]["image"]):
                s[side]["image"] = os.path.join(base, s[side]["image"])
        for v in (s.get("replace") or {}).values():
            if isinstance(v, dict) and "image" in v and not os.path.isabs(v["image"]):
                v["image"] = os.path.join(base, v["image"])
        if "from_slide" in s:
            build_cloned_slide(prs, originals, s, ea_font, i)
        else:
            build_layout_slide(prs, s, ea_font, i)

    if originals and not (args.keep_template_slides or spec.get("keep_template_slides")):
        sld_ids = prs.slides._sldIdLst
        for sld in list(sld_ids)[: len(originals)]:
            prs.part.drop_rel(sld.rId)
            sld_ids.remove(sld)
    prs.save(args.out)
    print(f"저장: {args.out} ({len(prs.slides)}장, 한글 글꼴: {ea_font or '테마 기본'})")
    if warnings:
        print("[경고]")
        for w in warnings:
            print("  ! " + w)
    print("다음: export_pdf.py로 PDF를 만들고 preview.py로 장별 렌더를 확인한다.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
