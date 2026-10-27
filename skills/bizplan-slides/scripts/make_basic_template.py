#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""사용자 템플릿이 없을 때 쓰는 기본 16:9 한국어 발표 템플릿(bizplan-basic.pptx)을 만든다.

    python3 -B make_basic_template.py [--out ../assets/templates/bizplan-basic.pptx] [--accent 1E3A8A]
흰 배경 + 남색 1색 + 강조 1색, 제목 32pt·본문 24/20/18pt, 한글 글꼴은 테마에 '맑은 고딕'으로 지정한다.
"""
import argparse
import copy
import os
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = "맑은 고딕"


def _theme(prs, accent, accent2):
    for rel in prs.slide_master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            part = rel.target_part
            xml = part.blob.decode("utf-8")
            xml = re.sub(r'(<a:latin typeface=")[^"]*(")', rf"\g<1>{FONT}\g<2>", xml)
            xml = re.sub(r'(<a:ea typeface=")[^"]*(")', rf"\g<1>{FONT}\g<2>", xml)
            xml = re.sub(r'(<a:accent1>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}', rf"\g<1>{accent}", xml)
            xml = re.sub(r'(<a:accent2>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}', rf"\g<1>{accent2}", xml)
            xml = re.sub(r'(<a:dk2>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}', rf"\g<1>{accent}", xml)
            part._blob = xml.encode("utf-8")


def _styles(prs, accent):
    st = prs.slide_master._element.find(qn("p:txStyles"))
    sizes = {"p:titleStyle": [3200], "p:bodyStyle": [2400, 2000, 1800, 1600, 1600]}
    for key, szs in sizes.items():
        style = st.find(qn(key))
        for i, sz in enumerate(szs, 1):
            lvl = style.find(qn(f"a:lvl{i}pPr"))
            if lvl is None:
                continue
            d = lvl.find(qn("a:defRPr"))
            d.set("sz", str(sz))
            if key == "p:titleStyle":
                d.set("b", "1")
                fill = d.find(qn("a:solidFill"))
                if fill is not None:
                    d.remove(fill)
                fill = d.makeelement(qn("a:solidFill"), {})
                clr = fill.makeelement(qn("a:srgbClr"), {"val": accent})
                fill.append(clr)
                d.insert(0, fill)
    # 제목 왼쪽 정렬(결론형 헤드라인)
    tl = st.find(qn("p:titleStyle")).find(qn("a:lvl1pPr"))
    tl.set("algn", "l")


def _layout_geometry(prs):
    """기본(4:3) 자리표시자를 16:9 화면에 맞게 다시 놓는다."""
    sw, sh = prs.slide_width, prs.slide_height
    m = Inches(0.6)
    title_box = (m, Inches(0.45), sw - 2 * m, Inches(1.0))
    body_box = (m, Inches(1.6), sw - 2 * m, sh - Inches(1.6) - Inches(0.8))
    for ph in prs.slide_master.placeholders:
        t = ph.placeholder_format.type
        name = str(t)
        if "TITLE" in name:
            ph.left, ph.top, ph.width, ph.height = title_box
        elif "BODY" in name:
            ph.left, ph.top, ph.width, ph.height = body_box
        elif "DATE" in name:
            ph.left, ph.top, ph.width, ph.height = m, sh - Inches(0.5), Inches(2.5), Inches(0.35)
        elif "FOOTER" in name:
            ph.left, ph.top, ph.width, ph.height = Inches(4.2), sh - Inches(0.5), sw - Inches(8.4), Inches(0.35)
        elif "SLIDE_NUMBER" in name:
            ph.left, ph.top, ph.width, ph.height = sw - m - Inches(1.5), sh - Inches(0.5), Inches(1.5), Inches(0.35)
    for layout in prs.slide_layouts:
        for ph in layout.placeholders:
            # 레이아웃마다 다른 글자 크기(예: 2단 28pt)를 지워 마스터 크기로 통일한다
            for d in ph._element.iter(qn("a:defRPr")):
                d.attrib.pop("sz", None)
            if ph._element.spPr.find(qn("a:xfrm")) is None:
                continue
            # 레이아웃 고유 위치는 10in(4:3) 기준이므로 가로를 비율대로 늘린다
            f = sw / Inches(10)
            ph.left, ph.width = int(ph.left * f), int(ph.width * f)


def _title_slide(prs):
    """표지 레이아웃은 제목·부제를 가운데 정렬로 되돌린다(마스터 제목은 왼쪽 정렬)."""
    for ph in prs.slide_layouts[0].placeholders:
        if ph.placeholder_format.idx not in (0, 1):
            continue
        lst = ph._element.txBody.find(qn("a:lstStyle"))
        lvl = lst.find(qn("a:lvl1pPr"))
        if lvl is None:
            lvl = lst.makeelement(qn("a:lvl1pPr"), {})
            lst.append(lvl)
        lvl.set("algn", "ctr")


def _master_bar(prs, accent):
    """마스터에 얇은 상단 강조선(모든 슬라이드에 나타남)."""
    tmp = prs.slides.add_slide(prs.slide_layouts[6])
    bar = tmp.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(accent)
    bar.line.fill.background()
    bar.name = "상단 강조선"
    prs.slide_master.shapes._spTree.append(copy.deepcopy(bar._element))
    sld = prs.slides._sldIdLst[0]
    prs.part.drop_rel(sld.rId)
    prs.slides._sldIdLst.remove(sld)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=os.path.join(HERE, "..", "assets", "templates", "bizplan-basic.pptx"))
    ap.add_argument("--accent", default="1E3A8A", help="기본색(남색)")
    ap.add_argument("--accent2", default="F59E0B", help="강조색")
    args = ap.parse_args()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    _theme(prs, args.accent, args.accent2)
    _styles(prs, args.accent)
    _layout_geometry(prs)
    _title_slide(prs)
    _master_bar(prs, args.accent)
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    prs.save(args.out)
    print(f"저장: {os.path.abspath(args.out)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
