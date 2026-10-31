#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""사업계획서 .md의 표·본문 숫자를 검산한다 — 심사위원의 암산 한 번에 걸리는 오류를 먼저 찾는다.

    python3 -B check_tables.py 사업계획서.md [--xlsx 검산표.xlsx] [--tolerance 0.5]
확인 항목
  1. 합계 행(합계·총계·계·Total)이 위 행들의 합과 같은가 — 열마다, 단위(원·천원·만원·백만원·억·조) 환산 후 비교
  2. 비율·비중(%) 열의 합이 100%인가
  3. 본문의 곱셈식 "A × B = C"(예: 1.5조 × 0.1% = 15억)가 맞는가
--xlsx를 주면 표마다 시트를 만들어 합계 행을 =SUM 수식으로 넣는다(제출 전 검산·수정용).
종료 코드: 오류가 있으면 1. 단위가 섞였거나 읽을 수 없는 칸은 '판단 보류'로 따로 보고한다.
"""
import argparse
import re
import sys

UNITS = [("천만원", 1e7), ("백만원", 1e6), ("천만", 1e7), ("백만", 1e6), ("조원", 1e12), ("억원", 1e8),
         ("만원", 1e4), ("천원", 1e3), ("조", 1e12), ("억", 1e8), ("만", 1e4), ("천", 1e3), ("원", 1)]
TOTAL = re.compile(r"^\s*(\*\*)?\s*(합\s*계|총\s*계|총\s*합|총\s*액|총\s*사업비|소\s*계|계|total|sum)\b", re.I)
SUBTOTAL = re.compile(r"소\s*계")
PLACEHOLDER = re.compile(r"○○|\[(?:작성 필요|산정 예정|공고 확인 필요|미확보)")
NUM = r"[0-9]+(?:,[0-9]{3})*(?:\.[0-9]+)?"
COMPOUND = re.compile(rf"({NUM})\s*(조|억|천만|백만|만|천)?\s*(?:({NUM})\s*(억|천만|백만|만|천))?\s*(원|%|명|개|건|회|가구|곳|대|톤|t|㎡)?")


def parse_amount(text, header_unit=None):
    """'2억 8,000만원', '20,000,000원', '40%', '1.2조'를 (값, 종류) 로. 못 읽으면 None."""
    t = text.replace("**", "").replace(" ", "").strip()
    if not t or t in ("-", "–", "—"):
        return None
    t = re.sub(r"\(.*?\)|\[.*?\]", "", t)
    m = re.fullmatch(rf"(?:약|~|±)?({NUM})%", t)
    if m:
        return float(m.group(1).replace(",", "")), "%"
    total, matched, kind = 0.0, False, None
    rest = t
    for m in re.finditer(rf"({NUM})(조원|억원|천만원|백만원|만원|천원|조|억|천만|백만|만|천|원)?", t):
        val = float(m.group(1).replace(",", ""))
        unit = m.group(2)
        mult = dict(UNITS).get(unit, None) if unit else None
        if mult is None:
            mult = header_unit or 1
            kind = kind or ("money" if header_unit else "num")
        else:
            kind = "money"
        total += val * mult
        matched = True
        rest = rest.replace(m.group(0), "", 1)
    count = re.search(r"(명|개|건|회|가구|곳|대|톤|t|㎡)$", t)
    rest = re.sub(r"^(약|~)|[원명개건회가구곳대톤t㎡]+$", "", rest)
    if not matched or rest.strip(" ,.~"):
        return None
    return total, "num" if count else kind


def header_unit(h):
    m = re.search(r"[(\[]\s*(조원|억원|천만원|백만원|만원|천원|원|억|만|조)\s*[)\]]|단위\s*[:：]\s*(조원|억원|천만원|백만원|만원|천원|원)", h)
    if not m:
        return None
    u = m.group(1) or m.group(2)
    return dict(UNITS)[u]


def tables(lines):
    i = 0
    while i < len(lines) - 1:
        if lines[i].strip().startswith("|") and re.match(r"^\|?\s*:?-{2,}", lines[i + 1].strip()):
            start = i
            header = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append((i + 1, [c.strip() for c in lines[i].strip().strip("|").split("|")]))
                i += 1
            yield start + 1, header, rows
        else:
            i += 1


def fmt(v, kind):
    if kind == "%":
        return f"{v:g}%"
    if kind == "money":
        if abs(v) >= 1e8:
            return f"{v / 1e8:,.2f}억원".replace(".00억", "억")
        if abs(v) >= 1e4:
            return f"{v / 1e4:,.0f}만원"
        return f"{v:,.0f}원"
    return f"{v:,.2f}".rstrip("0").rstrip(".")


def check_table(line_no, header, rows, tol, errors, holds):
    if not rows:
        return
    ncols = len(header)
    totals = [(ln, r) for ln, r in rows if r and TOTAL.match(r[0].replace("*", "")) and not SUBTOTAL.search(r[0])]
    body = [(ln, r) for ln, r in rows if not (r and TOTAL.match(r[0].replace("*", "")))]
    for c in range(1, ncols):
        hu = header_unit(header[c]) or header_unit(" ".join(header))
        is_pct_col = "%" in header[c] or re.search(r"비율|비중|구성비|점유율", header[c])
        vals, blank = [], 0
        for ln, r in body:
            if c < len(r):
                if PLACEHOLDER.search(r[c]):
                    blank += 1          # 미확보 칸 — 숫자가 아니므로 합계를 검산할 수 없다
                    continue
                v = parse_amount(r[c], hu)
                if v:
                    vals.append((ln, v))
        summed = any(c < len(r) and parse_amount(r[c], hu) for _, r in totals)   # 합계 행에 숫자가 있는 열만(글 열 '산출근거' 등은 제외)
        if blank and totals and summed:
            holds.append(f"{totals[0][0]}행 '{header[c]}' 열: 미확보 값(○○) {blank}칸 — 합계를 검산하지 못해 판단 보류")
            continue
        if not vals:
            continue
        kinds = {k for _, (_, k) in vals}
        for ln, r in totals:
            if c >= len(r):
                continue
            stated = parse_amount(r[c], hu)
            if not stated:
                continue
            same = [v for _, (v, k) in vals if k == stated[1]]
            if len(kinds) > 1 or len(same) != len(vals):
                holds.append(f"{ln}행 '{header[c]}' 열: 단위가 섞여 합계를 판단 보류({', '.join(sorted(kinds))})")
                continue
            if len(vals) < 2:
                continue
            calc = sum(same)
            if stated[1] == "%" or is_pct_col:
                if abs(calc - stated[0]) > 1.0:
                    errors.append(f"{ln}행 '{header[c]}' 열 합계: 표기 {fmt(stated[0], '%')} ≠ 위 행 합 {fmt(calc, '%')}")
            elif abs(calc - stated[0]) > max(abs(stated[0]) * tol / 100, 0.5):
                errors.append(f"{ln}행 '{header[c]}' 열 합계: 표기 {fmt(*stated)} ≠ 위 행 합 {fmt(calc, stated[1])}"
                              f" (차이 {fmt(stated[0] - calc, stated[1])})")
        if is_pct_col and not totals and kinds == {"%"} and len(vals) >= 2:
            s = sum(v for _, (v, _) in vals)
            if abs(s - 100) > 1.0 and s > 50:
                errors.append(f"{line_no}행 표 '{header[c]}' 열: 비율 합이 {s:g}%다(100%가 아님)")


def check_formulas(lines, tol, errors):
    """본문의 'A × B (× C) = D' 꼴을 검산한다. %는 비율로, 조·억·만은 원 단위로 환산한다."""
    pat = re.compile(rf"((?:{NUM}\s*(?:조|억|천만|백만|만|천)?\s*(?:원|%|명|개|건|회|가구|곳|대)?\s*[×xX\*]\s*)+"
                     rf"{NUM}\s*(?:조|억|천만|백만|만|천)?\s*(?:원|%|명|개|건|회|가구|곳|대)?)\s*(?:=|≈|→)\s*"
                     rf"({NUM}\s*(?:조|억|천만|백만|만|천)?\s*(?:원|%|명|개|건|회)?)")
    for no, line in enumerate(lines, 1):
        if line.strip().startswith("|") and "=" not in line:
            continue
        for m in pat.finditer(line):
            factors = re.split(r"\s*[×xX\*]\s*", m.group(1).strip())
            prod, money = 1.0, False
            ok = True
            for f in factors:
                pa = parse_amount(f)
                if not pa:
                    ok = False
                    break
                v, k = pa
                prod *= v / 100 if k == "%" else v
                money = money or k == "money"
            res = parse_amount(m.group(2))
            if not ok or not res:
                continue
            rv = res[0] / 100 if res[1] == "%" else res[0]
            if rv and abs(prod - rv) > abs(rv) * max(tol, 2) / 100:
                errors.append(f"{no}행 식 '{m.group(0).strip()}': 계산하면 {fmt(prod, 'money' if money else 'num')}")


def to_xlsx(path, tbls):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for n, (line_no, header, rows) in enumerate(tbls, 1):
        ws = wb.create_sheet(f"표{n}_{line_no}행")
        ws.append(header)
        body_rows = []
        for ln, r in rows:
            is_total = bool(r and TOTAL.match(r[0].replace("*", "")))
            out = []
            for c, cell in enumerate(r):
                pa = parse_amount(cell, header_unit(header[c]) if c < len(header) else None) if c else None
                out.append(pa[0] / 100 if pa and pa[1] == "%" else pa[0] if pa else cell.replace("**", ""))
            ws.append(out)
            row_idx = ws.max_row
            if is_total and body_rows:
                for c in range(2, len(r) + 1):
                    col = get_column_letter(c)
                    if all(isinstance(ws[f"{col}{b}"].value, (int, float)) for b in body_rows):
                        ws[f"{col}{row_idx}"] = f"=SUM({col}{body_rows[0]}:{col}{body_rows[-1]})"
            elif not is_total:
                body_rows.append(row_idx)
        for c in range(1, len(header) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 18
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("md")
    ap.add_argument("--xlsx")
    ap.add_argument("--tolerance", type=float, default=0.5, help="합계 허용 오차(%%)")
    args = ap.parse_args()
    with open(args.md, encoding="utf-8") as f:
        lines = f.read().splitlines()
    tbls = list(tables(lines))
    errors, holds = [], []
    blanks = sum(1 for _, _, rows in tbls for _, r in rows for cell in r if PLACEHOLDER.search(cell))
    for line_no, header, rows in tbls:
        check_table(line_no, header, rows, args.tolerance, errors, holds)
    check_formulas(lines, args.tolerance, errors)
    print(f"표 {len(tbls)}개 검산")
    for e in errors:
        print("  ✕ " + e)
    for h in holds:
        print("  △ " + h)
    if not errors:
        print("  합계·비율·곱셈식 오류 없음" + (" (판단 보류는 직접 확인)" if holds else ""))
    if blanks:
        print(f"  미확보 칸(○○·[작성 필요]) {blanks}개는 검산하지 않았다 — 숫자가 들어오면 다시 돌린다")
    if args.xlsx:
        to_xlsx(args.xlsx, tbls)
        print(f"검산용 엑셀: {args.xlsx} (합계 행은 =SUM 수식)")
    print("한계: 표 밖 서술 숫자 간 모순(예: 같은 매출이 쪽마다 다른 단위로 적힘)과 CAGR 서술은 직접 대조한다.")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
