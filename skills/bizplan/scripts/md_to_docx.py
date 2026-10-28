#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""누적한 사업계획서 .md를 제출 서식에 가까운 .docx로 바꾼다.

    python3 -B md_to_docx.py 사업계획서.md 사업계획서.docx [--gov-bullets] [--font "맑은 고딕"]
                             [--resource-path 그림폴더] [--engine auto|pandoc|python] [--captions]
서식: A4·여백 20mm, 한글 글꼴(기본 맑은 고딕)·ko-KR, 검정 제목, 표 격자 테두리·머리행 굵게·열 너비 내용 비례,
그림은 본문 폭에 맞추고 앞 문단(헤드라인)과 같은 쪽에 묶는다(alt 캡션은 --captions일 때만). --gov-bullets는 불릿 계층을 공문서 개조식(□ → ○ → - → ·)으로 바꾼다.
pandoc이 있으면 기준 서식(reference.docx)을 만들어 쓰고, 없으면 python-docx로 직접 변환한다.
"""
import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

GOV = ["□", "○", "-", "·"]


# ---------------------------------------------------------------- 전처리

def gov_bullets(md):
    """마크다운 불릿을 □○-· 접두 문단으로 바꾼다(표·코드 블록 안은 그대로)."""
    out, fence = [], False
    indents = [len(m.group(1)) for m in re.finditer(r"^( *)[-*+] ", md, re.M) if m.group(1)]
    unit = min(indents) if indents else 2
    for line in md.splitlines():
        if line.lstrip().startswith("```"):
            fence = not fence
        m = re.match(r"^( *)[-*+] (.*)$", line)
        if m and not fence and not line.lstrip().startswith("|"):
            level = min(len(m.group(1)) // unit, len(GOV) - 1)
            text = m.group(2)
            # 글이 이미 ○·□ 같은 개조식 기호로 시작하면 기호를 겹쳐 붙이지 않는다('- ○ [한 줄 소개]' → '○ [한 줄 소개]')
            mark = "" if re.match(r"^[□○◦·■●▪-] ", text) else f"{GOV[level]} "
            # 문단으로 쓰되 들여쓰기는 전각 공백으로 표현(목록 번호 체계 없이 공문서처럼)
            out.append("")
            out.append("　" * level + f"{mark}{text}")
            out.append("")
        else:
            out.append(line)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out))


# ---------------------------------------------------------------- 서식 공통

def _set_fonts(doc, font):
    for st in doc.styles:
        try:
            f = st.font
        except AttributeError:
            continue
        if f is None:
            continue
        f.name = font
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), font)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(a), None)
        if st.name.startswith("Heading") or st.name == "Title":
            f.color.rgb = RGBColor(0, 0, 0)
    for lang in doc.styles.element.iter(qn("w:lang")):
        lang.set(qn("w:val"), "ko-KR")
        lang.set(qn("w:eastAsia"), "ko-KR")


def _page(doc):
    for s in doc.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin = s.right_margin = Mm(20)
        s.top_margin = s.bottom_margin = Mm(20)


def _borders(tbl):
    tblPr = tbl._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement(f"w:{e}")
        x.set(qn("w:val"), "single")
        x.set(qn("w:sz"), "4")
        x.set(qn("w:color"), "000000")
        b.append(x)
    tblPr.append(b)


def _em(text):
    # 한글·한자와 ①○★□ 같은 기호(전각으로 찍힌다)는 1em, 영문 대문자·%는 0.85em, 나머지 영문·숫자는 0.62em
    return sum(0.33 if c.isspace() else 1.0 if ord(c) > 0x2E7F or 0x2460 <= ord(c) <= 0x27BF
               else 0.85 if "A" <= c <= "Z" or c in "%@" else 0.62 for c in text)


def _grid(tbl, widths):
    """열 너비를 칸(tcW)과 표 격자(tblGrid) 둘 다에 쓴다 — LibreOffice·한글은 격자 값을 따른다."""
    tbl.autofit = False
    grid = tbl._tbl.tblGrid
    cols = grid.findall(qn("w:gridCol")) if grid is not None else []
    for c, w in enumerate(widths):
        if c < len(cols):
            cols[c].set(qn("w:w"), str(int(w.twips)))
        for row in tbl.rows:
            if c < len(row.cells):
                row.cells[c].width = w


def _widths(tbl, total_mm=170):
    """열 너비를 열별 최장 글 길이에 비례(최소 10%, 최대 55%)로 정하되, 열마다 가장 긴 낱말은 끊기지 않게 한다.
    7열 이상인 넓은 표(경쟁 비교 등)는 글자를 9pt로 줄인다."""
    ncols = len(tbl.columns)
    if not ncols:
        return
    wide = ncols >= 7
    em_mm = 3.3 if wide else 4.3                      # 한 글자 폭: 넓은 표 9pt · 기본 12pt(pandoc 기준 서식) + 여유
    lens, mins = [], []
    for c in range(ncols):
        texts = [row.cells[c].text for row in tbl.rows if c < len(row.cells)]
        lens.append(max([_em(t) for t in texts] + [2]))
        # 끊기면 읽기 어려운 낱말 폭: 한글 낱말은 4자, 영문·숫자 낱말(제품명 등)은 6em까지만 지킨다(URL은 끊겨도 된다)
        word = max([min(_em(w), 4 if re.search("[가-힣]", w) else 6) for t in texts for w in t.split()] + [1])
        mins.append(word * em_mm + 4.6)                # + 칸 안 좌우 여백(LibreOffice·한글 기본 약 1.9mm × 2)과 여유
    shares = [min(max(l / sum(lens), 0.10), 0.55) for l in lens]
    w = [total_mm * x / sum(shares) for x in shares]
    if sum(mins) >= total_mm:                          # 최소 폭만으로 넘치면 최소 폭 비율로 나눈다
        w = [total_mm * m / sum(mins) for m in mins]
    else:
        for _ in range(4):                             # 모자란 열을 최소 폭까지 올리고, 나머지 열에서 그만큼 덜어 낸다
            low = [c for c in range(ncols) if w[c] < mins[c]]
            if not low:
                break
            for c in low:
                w[c] = mins[c]
            free = [c for c in range(ncols) if c not in low]
            rest = total_mm - sum(w[c] for c in low)
            tot = sum(w[c] for c in free) or 1
            for c in free:
                w[c] = w[c] * rest / tot
    _grid(tbl, [Mm(x) for x in w])
    if wide:
        for row in tbl.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    for r in par.runs:
                        r.font.size = Pt(9)


def _image_grid(tbl, total_mm=170, max_h=Mm(60)):
    """그림이 든 표(증빙 사진 격자)는 칸 너비를 같게 하고, 그림을 칸 안(높이 60mm 이하)에 맞춘다."""
    from docx.shape import InlineShape
    ncols = len(tbl.columns) or 1
    col = Mm(total_mm / ncols)
    _grid(tbl, [col] * ncols)
    for inline in tbl._tbl.xpath(".//wp:inline"):
        shape = InlineShape(inline)
        for limit, side in ((col - Mm(4), "width"), (max_h, "height")):
            if getattr(shape, side) > limit:
                ratio = limit / getattr(shape, side)
                shape.width, shape.height = int(shape.width * ratio), int(shape.height * ratio)


def _shade(p, fill="FFF8E6", bar="F59E0B"):
    """인용 문단에 옅은 바탕과 왼쪽 막대를 넣는다(작성 가이드의 '이렇게 써 보세요' 예문이 설명과 구분되게)."""
    pPr = p._p.get_or_add_pPr()
    for tag in ("w:pBdr", "w:shd"):
        old = pPr.find(qn(tag))
        if old is not None:
            pPr.remove(old)
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "6"), ("w:color", bar)):
        left.set(qn(k), v)
    bdr.append(left)
    shd = OxmlElement("w:shd")
    for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
        shd.set(qn(k), v)
    # pPr 자식 순서: pStyle, keepNext, …, numPr, …, pBdr, shd, … — numPr 뒤에 둔다
    anchor = None
    for tag in ("w:numPr", "w:suppressLineNumbers", "w:keepLines", "w:keepNext", "w:pStyle"):
        anchor = pPr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addnext(bdr)
    else:
        pPr.insert(0, bdr)
    bdr.addnext(shd)


def _keep_with_next(doc):
    """그림·표 바로 앞 문단(헤드라인·표 제목)이 쪽 끝에 홀로 남지 않게 다음 것과 같은 쪽에 묶는다."""
    prev = None
    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "tbl" or (tag == "p" and el.xpath(".//w:drawing")):
            if prev is not None:
                prev.get_or_add_pPr().keepNext_val = True
        if tag == "p" and "".join(el.itertext()).strip():
            prev = el
        elif tag == "tbl":
            prev = None


def postprocess(path, font):
    doc = Document(path)
    _page(doc)
    _set_fonts(doc, font)
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if p.style.name == "Block Text":
            _shade(p)
            if p.text.strip().startswith("디벨롭 가이드"):     # 상자 제목 뒤 목록(Compact)까지 같은 상자로
                j = i + 1
                while j < len(paras) and paras[j].style.name in ("Compact", "Block Text"):
                    _shade(paras[j])
                    j += 1
    for tbl in doc.tables:
        _borders(tbl)
        for row in tbl.rows:                  # 한 행이 두 쪽에 걸쳐 갈라지지 않게(합계 행이 '2,000 / 만 원'으로 쪼개지던 문제)
            trPr = row._tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:
                cs = OxmlElement("w:cantSplit")       # 스키마 순서: cantSplit은 trHeight·tblHeader보다 앞
                nxt = next((trPr.find(qn(t)) for t in ("w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden")
                            if trPr.find(qn(t)) is not None), None)
                nxt.addprevious(cs) if nxt is not None else trPr.append(cs)
        if tbl._tbl.xpath(".//w:drawing"):
            _image_grid(tbl)                  # 사진 격자: 같은 너비 칸, 그림은 칸에 맞춘다
        else:
            _widths(tbl)
        if tbl.rows:
            for cell in tbl.rows[0].cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True
    _keep_with_next(doc)
    # 그림이 본문 폭(170mm)을 넘지 않게, 정사각형 도식이 한 쪽을 다 차지하지 않게 높이는 120mm까지
    for shape in doc.inline_shapes:
        if shape._inline.xpath("ancestor::w:tc"):
            continue                          # 표 안 그림은 _image_grid가 맞춘다
        for limit, side in ((Mm(170), "width"), (Mm(120), "height")):
            if getattr(shape, side) > limit:
                ratio = limit / getattr(shape, side)
                shape.width, shape.height = int(shape.width * ratio), int(shape.height * ratio)
    doc.save(path)


# ---------------------------------------------------------------- pandoc 경로

def reference_docx(path, font):
    data = subprocess.run(["pandoc", "-o", path, "--print-default-data-file", "reference.docx"],
                          capture_output=True)
    if data.returncode != 0:
        raise RuntimeError(data.stderr.decode("utf-8", "ignore"))
    doc = Document(path)
    _page(doc)
    _set_fonts(doc, font)
    doc.save(path)


def via_pandoc(md_text, out, font, resource_path, captions=False):
    with tempfile.TemporaryDirectory() as tmp:
        ref = os.path.join(tmp, "reference.docx")
        reference_docx(ref, font)
        src = os.path.join(tmp, "in.md")
        with open(src, "w", encoding="utf-8") as f:
            f.write(re.sub(r"<u>([^<]+)</u>", r"[\1]{.underline}", md_text))   # pandoc은 <u>를 docx에서 버린다
        # 시각자료는 그림 안에 결론 제목이 있다 — alt 글을 그림 아래 캡션으로 한 번 더 달지 않는다(--captions로 켠다)
        cmd = ["pandoc", src, "-f", "markdown" if captions else "markdown-implicit_figures", "--reference-doc", ref, "-o", out,
               f"--resource-path={resource_path}"]
        r = subprocess.run(cmd, capture_output=True, text=True)
        if r.returncode != 0:
            raise RuntimeError(r.stderr)
        return [l for l in r.stderr.splitlines() if "WARNING" in l]


# ---------------------------------------------------------------- python-docx 경로

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\{\.underline\}|<u>[^<]+</u>)")


def _runs(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        elif part.endswith("{.underline}"):
            par.add_run(part[1:part.index("]{")]).underline = True
        elif part.startswith("<u>"):
            par.add_run(part[3:-4]).underline = True
        else:
            par.add_run(part)


def _cells(line):
    # 표 칸의 그림 표기는 python-docx 경로에서 '[그림: 설명]'으로 남긴다(pandoc 경로는 실제 그림을 넣는다)
    return [re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"[그림: \1]", c.strip()) for c in line.strip().strip("|").split("|")]


def via_python(md_text, out, resource_path, captions=False):
    doc = Document()
    lines = md_text.splitlines()
    i, warnings = 0, []
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s or re.match(r"^-{3,}$|^\*{3,}$", s):
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(buf))
            r.font.name, r.font.size = "Consolas", Pt(9)
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            doc.add_heading(m.group(2).replace("**", ""), level=min(len(m.group(1)), 4))
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-{2,}", lines[i + 1].strip()):
            header = _cells(s)
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_cells(lines[i]))
                i += 1
            t = doc.add_table(rows=1 + len(rows), cols=len(header))
            for ci, h in enumerate(header):
                _runs(t.cell(0, ci).paragraphs[0], h)
            for ri, row in enumerate(rows, 1):
                for ci in range(len(header)):
                    _runs(t.cell(ri, ci).paragraphs[0], row[ci] if ci < len(row) else "")
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s)
        if m:
            path = m.group(2) if os.path.isabs(m.group(2)) else os.path.join(resource_path, m.group(2))
            if os.path.exists(path):
                doc.add_picture(path, width=Mm(160))
                if m.group(1) and captions:
                    doc.add_paragraph(m.group(1))
            else:
                warnings.append(f"그림 없음: {m.group(2)}")
                doc.add_paragraph(f"[그림: {m.group(1)}]")
            i += 1
            continue
        m = re.match(r"^( *)([-*+]|\d+[.)])\s+(.*)$", line)
        if m:
            level = len(m.group(1)) // 2
            style = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
            if level:
                style += f" {min(level + 1, 3)}"
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph()
            _runs(p, m.group(3))
            i += 1
            continue
        if s.startswith(">"):
            body = s.lstrip(">").strip()
            if not body:
                i += 1
                continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(6)
            _runs(p, body)
            _shade(p)
            i += 1
            continue
        buf = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||!\[|```|>|( *)([-*+]|\d+[.)])\s)", lines[i].strip()):
            buf.append(lines[i].strip())
            i += 1
        _runs(doc.add_paragraph(), " ".join(buf))
    doc.save(out)
    return warnings


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("md")
    ap.add_argument("out")
    ap.add_argument("--gov-bullets", action="store_true")
    ap.add_argument("--font", default="맑은 고딕")
    ap.add_argument("--resource-path", help="그림 상대 경로의 기준(기본: .md가 있는 폴더)")
    ap.add_argument("--engine", choices=("auto", "pandoc", "python"), default="auto")
    ap.add_argument("--captions", action="store_true", help="그림 alt 글을 그림 아래 캡션으로도 단다(기본: 달지 않음)")
    args = ap.parse_args()
    with open(args.md, encoding="utf-8") as f:
        md = f.read()
    if args.gov_bullets:
        md = gov_bullets(md)
    res = args.resource_path or os.path.dirname(os.path.abspath(args.md))
    engine = args.engine
    if engine == "auto":
        engine = "pandoc" if shutil.which("pandoc") else "python"
    warnings = (via_pandoc(md, args.out, args.font, res, args.captions) if engine == "pandoc"
                else via_python(md, args.out, res, args.captions))
    postprocess(args.out, args.font)
    doc = Document(args.out)
    print(f"저장: {args.out} (엔진: {engine}, 표 {len(doc.tables)}개, 그림 {len(doc.inline_shapes)}개)")
    for w in warnings:
        print("  ! " + w)
    print("검수: 한글(HWP)에서 열어 표 줄바꿈·빈 셀·그림 위치를 확인한다. 제출 양식이 HWP면 md_to_hwpx.py도 있다.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
